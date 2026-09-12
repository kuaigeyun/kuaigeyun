"""固定资产卡片服务。"""

from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook, load_workbook

from apps.common.audit_actor import apply_create_audit
from apps.kuaicaiwu.models.fixed_asset import FaAsset, FaCategory
from apps.kuaicaiwu.services.fa_core import (
    compute_monthly_depreciation,
    compute_net_value,
    generate_daily_code,
    model_to_dict,
    quantize_money,
    touch_updated,
)
from core.utils.timezone_utils import resolve_business_datetime, to_site_date
from infra.exceptions.exceptions import BusinessLogicError, NotFoundError, ValidationError
from infra.models.user import User


def _parse_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip()[:10])


class FaAssetService:
    async def _resolve_category_name(
        self, tenant_id: int, payload: dict[str, Any]
    ) -> dict[str, Any]:
        """资产类别仅作分类；折旧参数在卡片上逐台维护，不从类别继承。"""
        category_id = payload.get("category_id")
        if not category_id:
            return payload
        cat = await FaCategory.get_or_none(
            id=int(category_id), tenant_id=tenant_id, deleted_at__isnull=True
        )
        if not cat:
            raise ValidationError("资产类别不存在")
        payload["category_name"] = cat.category_name
        return payload

    def _recalc_depreciation_fields(self, payload: dict[str, Any]) -> dict[str, Any]:
        original = quantize_money(payload.get("original_value") or 0)
        residual_rate = quantize_money(payload.get("residual_rate") or "0.05")
        useful_life = int(payload.get("useful_life_months") or 60)
        accumulated = quantize_money(payload.get("accumulated_depreciation") or 0)
        impairment = quantize_money(payload.get("impairment_value") or 0)
        monthly = compute_monthly_depreciation(original, residual_rate, useful_life)
        payload["monthly_depreciation"] = monthly
        payload["net_value"] = float(compute_net_value(original, accumulated, impairment))
        return payload

    async def list_assets(
        self,
        tenant_id: int,
        *,
        skip: int = 0,
        limit: int = 50,
        keyword: Optional[str] = None,
        status: Optional[str] = None,
        category_id: Optional[int] = None,
    ) -> dict[str, Any]:
        q = FaAsset.filter(tenant_id=tenant_id, deleted_at__isnull=True)
        if status:
            q = q.filter(status=status)
        if category_id:
            q = q.filter(category_id=category_id)
        if keyword and keyword.strip():
            from tortoise.expressions import Q

            term = keyword.strip()
            q = q.filter(Q(asset_code__icontains=term) | Q(asset_name__icontains=term))
        total = await q.count()
        rows = await q.order_by("-created_at", "-id").offset(skip).limit(limit)
        items = []
        for row in rows:
            item = model_to_dict(row)
            item["net_value"] = float(
                compute_net_value(
                    quantize_money(row.original_value),
                    quantize_money(row.accumulated_depreciation),
                    quantize_money(row.impairment_value),
                )
            )
            items.append(item)
        return {"items": items, "total": total, "skip": skip, "limit": limit}

    async def get_asset(self, tenant_id: int, asset_id: int) -> dict[str, Any]:
        row = await FaAsset.get_or_none(id=asset_id, tenant_id=tenant_id, deleted_at__isnull=True)
        if not row:
            raise NotFoundError("资产不存在")
        item = model_to_dict(row)
        item["net_value"] = float(
            compute_net_value(
                quantize_money(row.original_value),
                quantize_money(row.accumulated_depreciation),
                quantize_money(row.impairment_value),
            )
        )
        return item

    async def _resolve_category_id(
        self, tenant_id: int, *, category_id: Any = None, category_name: Any = None
    ) -> Optional[int]:
        if category_id:
            return int(category_id)
        name = str(category_name or "").strip()
        if not name:
            return None
        cat = await FaCategory.filter(
            tenant_id=tenant_id,
            category_name=name,
            deleted_at__isnull=True,
            is_active=True,
        ).first()
        if not cat:
            cat = await FaCategory.filter(
                tenant_id=tenant_id,
                category_code=name,
                deleted_at__isnull=True,
                is_active=True,
            ).first()
        if not cat:
            raise ValidationError(f"资产类别不存在: {name}")
        return int(cat.id)

    async def create_asset(
        self, tenant_id: int, data: dict[str, Any], user: User | int
    ) -> dict[str, Any]:
        resolved_user = user if isinstance(user, User) else await User.get(id=int(user))
        category_id = await self._resolve_category_id(
            tenant_id,
            category_id=data.get("category_id"),
            category_name=data.get("category_name"),
        )
        asset_code = (data.get("asset_code") or "").strip()
        if not asset_code:
            asset_code = await generate_daily_code(FaAsset, tenant_id, "FA", "asset_code")
        asset_name = (data.get("asset_name") or "").strip()
        if not asset_name:
            raise ValidationError("资产名称必填")
        exists = await FaAsset.filter(
            tenant_id=tenant_id, asset_code=asset_code, deleted_at__isnull=True
        ).exists()
        if exists:
            raise ValidationError(f"资产编号 {asset_code} 已存在")

        payload: dict[str, Any] = {
            "tenant_id": tenant_id,
            "asset_code": asset_code,
            "asset_name": asset_name,
            "category_id": category_id,
            "quantity": quantize_money(data.get("quantity") or 1),
            "unit": data.get("unit"),
            "change_method": data.get("change_method"),
            "department_id": data.get("department_id"),
            "department_name": data.get("department_name"),
            "user_id": data.get("user_id"),
            "user_name": data.get("user_name"),
            "status": data.get("status") or "active",
            "location": data.get("location"),
            "start_use_date": _parse_date(data.get("start_use_date")),
            "entry_date": _parse_date(data.get("entry_date")),
            "specification": data.get("specification"),
            "notes": data.get("notes"),
            "attachment_uuids": list(data.get("attachment_uuids") or []),
            "depreciation_method": data.get("depreciation_method") or "straight_line",
            "original_value": quantize_money(data.get("original_value") or 0),
            "impairment_value": quantize_money(data.get("impairment_value") or 0),
            "useful_life_months": int(data.get("useful_life_months") or 60),
            "depreciated_periods": int(data.get("depreciated_periods") or 0),
            "accumulated_depreciation": quantize_money(data.get("accumulated_depreciation") or 0),
            "residual_rate": quantize_money(data.get("residual_rate") or "0.05"),
            "asset_account_code": data.get("asset_account_code") or "1601",
            "accumulated_depreciation_account_code": data.get(
                "accumulated_depreciation_account_code"
            )
            or "1602",
            "expense_account_code": data.get("expense_account_code") or "6602",
            "source_kuaioa_asset_id": data.get("source_kuaioa_asset_id"),
            "source_purchase_id": data.get("source_purchase_id"),
        }
        payload = await self._resolve_category_name(tenant_id, payload)
        payload["monthly_depreciation"] = compute_monthly_depreciation(
            payload["original_value"],
            payload["residual_rate"],
            payload["useful_life_months"],
        )
        apply_create_audit(payload, resolved_user)
        row = await FaAsset.create(**payload)
        return await self.get_asset(tenant_id, row.id)

    async def update_asset(
        self, tenant_id: int, asset_id: int, data: dict[str, Any], user: User
    ) -> dict[str, Any]:
        row = await FaAsset.get_or_none(id=asset_id, tenant_id=tenant_id, deleted_at__isnull=True)
        if not row:
            raise NotFoundError("资产不存在")
        if row.status == "disposed":
            raise BusinessLogicError("已清理资产不可编辑")

        scalar_fields = (
            "asset_name",
            "category_id",
            "category_name",
            "quantity",
            "unit",
            "change_method",
            "department_id",
            "department_name",
            "user_id",
            "user_name",
            "status",
            "location",
            "specification",
            "notes",
            "depreciation_method",
            "useful_life_months",
            "depreciated_periods",
            "asset_account_code",
            "accumulated_depreciation_account_code",
            "expense_account_code",
        )
        for key in scalar_fields:
            if key in data and data[key] is not None:
                if key == "quantity":
                    setattr(row, key, quantize_money(data[key]))
                elif key in {"useful_life_months", "depreciated_periods", "category_id", "department_id", "user_id"}:
                    setattr(row, key, int(data[key]))
                else:
                    setattr(row, key, data[key])
        for key in ("start_use_date", "entry_date"):
            if key in data:
                setattr(row, key, _parse_date(data[key]))
        for key in ("original_value", "impairment_value", "accumulated_depreciation", "residual_rate"):
            if key in data and data[key] is not None:
                setattr(row, key, quantize_money(data[key]))
        if "attachment_uuids" in data:
            row.attachment_uuids = list(data["attachment_uuids"] or [])

        if "category_id" in data and data["category_id"]:
            resolved = await self._resolve_category_name(
                tenant_id, {"category_id": data["category_id"]}
            )
            if "category_name" in resolved:
                row.category_name = resolved["category_name"]

        row.monthly_depreciation = compute_monthly_depreciation(
            quantize_money(row.original_value),
            quantize_money(row.residual_rate),
            int(row.useful_life_months or 0),
        )
        await touch_updated(row, user)
        await row.save()
        return await self.get_asset(tenant_id, asset_id)

    async def delete_asset(self, tenant_id: int, asset_id: int, user: User) -> None:
        row = await FaAsset.get_or_none(id=asset_id, tenant_id=tenant_id, deleted_at__isnull=True)
        if not row:
            raise NotFoundError("资产不存在")
        if row.status not in {"active", "idle"}:
            raise BusinessLogicError("仅在用或闲置资产可删除")
        row.deleted_at = resolve_business_datetime()
        await touch_updated(row, user)
        await row.save()

    async def create_from_purchase(
        self,
        tenant_id: int,
        *,
        purchase_id: int,
        title: str,
        asset_category: Optional[str],
        quantity: int,
        amount: Optional[Decimal],
        applicant_id: Optional[int],
        applicant_name: Optional[str],
        department_name: Optional[str],
        attachment_uuids: Optional[list[str]],
        user_id: int,
    ) -> dict[str, Any]:
        existing = await FaAsset.filter(
            tenant_id=tenant_id, source_purchase_id=purchase_id, deleted_at__isnull=True
        ).count()
        if existing > 0:
            raise BusinessLogicError("该采买申请已建卡")

        category_id = None
        if asset_category:
            cat = await FaCategory.filter(
                tenant_id=tenant_id,
                category_name=asset_category,
                deleted_at__isnull=True,
                is_active=True,
            ).first()
            if not cat:
                cat = await FaCategory.filter(
                    tenant_id=tenant_id,
                    category_code=asset_category,
                    deleted_at__isnull=True,
                    is_active=True,
                ).first()
            if cat:
                category_id = cat.id

        qty = max(1, int(quantity or 1))
        unit_amount = quantize_money(amount or 0)
        if qty > 1 and unit_amount > 0:
            unit_amount = quantize_money(unit_amount / Decimal(qty))

        created: list[dict[str, Any]] = []
        entry_date = to_site_date(resolve_business_datetime())
        for index in range(qty):
            asset_name = title if qty == 1 else f"{title}-{index + 1}"
            item = await self.create_asset(
                tenant_id,
                {
                    "asset_name": asset_name,
                    "category_id": category_id,
                    "original_value": unit_amount,
                    "entry_date": entry_date.isoformat(),
                    "start_use_date": entry_date.isoformat(),
                    "user_id": applicant_id,
                    "user_name": applicant_name,
                    "department_name": department_name,
                    "attachment_uuids": list(attachment_uuids or []),
                    "source_purchase_id": purchase_id,
                    "change_method": "purchase",
                },
                user_id,
            )
            created.append(item)
        return {"assets": created, "count": len(created), "asset": created[0] if created else None}

    async def migrate_from_kuaioa_asset(
        self, tenant_id: int, oa_asset: Any, user: User
    ) -> dict[str, Any]:
        existing = await FaAsset.filter(
            tenant_id=tenant_id,
            source_kuaioa_asset_id=oa_asset.id,
            deleted_at__isnull=True,
        ).first()
        if existing:
            return model_to_dict(existing)

        status_map = {
            "in_stock": "idle",
            "issued": "active",
            "finance_pending": "active",
            "written_off": "disposed",
            "scrapped": "scrapped",
        }
        category_id = None
        if oa_asset.asset_category:
            cat = await FaCategory.filter(
                tenant_id=tenant_id,
                category_name=oa_asset.asset_category,
                deleted_at__isnull=True,
            ).first()
            if cat:
                category_id = cat.id

        return await self.create_asset(
            tenant_id,
            {
                "asset_code": oa_asset.asset_code,
                "asset_name": oa_asset.asset_name,
                "category_id": category_id,
                "original_value": oa_asset.purchase_amount or 0,
                "entry_date": oa_asset.purchase_date.isoformat() if oa_asset.purchase_date else None,
                "start_use_date": oa_asset.purchase_date.isoformat() if oa_asset.purchase_date else None,
                "user_id": oa_asset.custodian_id,
                "user_name": oa_asset.custodian_name,
                "department_name": oa_asset.department_name,
                "location": oa_asset.location,
                "status": status_map.get(oa_asset.status or "active", "active"),
                "notes": oa_asset.notes,
                "attachment_uuids": list(oa_asset.attachment_uuids or []),
                "source_kuaioa_asset_id": oa_asset.id,
                "source_purchase_id": oa_asset.purchase_id,
                "change_method": "migrate",
            },
            user,
        )

    async def export_excel(self, tenant_id: int) -> BytesIO:
        from apps.kuaicaiwu.services.fa_asset_import_template import FA_ASSET_IMPORT_HEADERS

        result = await self.list_assets(tenant_id, skip=0, limit=10000)
        wb = Workbook()
        ws = wb.active
        ws.title = "资产清单"
        ws.append(FA_ASSET_IMPORT_HEADERS)
        status_labels = {
            "active": "在用",
            "idle": "闲置",
            "disposed": "已清理",
            "scrapped": "已报废",
        }
        depr_labels = {"straight_line": "年限平均法"}
        for item in result["items"]:
            ws.append(
                [
                    item.get("asset_code"),
                    item.get("asset_name"),
                    item.get("category_name"),
                    item.get("change_method"),
                    item.get("quantity"),
                    item.get("unit"),
                    item.get("useful_life_months"),
                    item.get("department_name"),
                    item.get("user_name"),
                    status_labels.get(str(item.get("status") or ""), item.get("status")),
                    item.get("location"),
                    item.get("start_use_date"),
                    item.get("entry_date"),
                    item.get("specification"),
                    item.get("notes"),
                    depr_labels.get(str(item.get("depreciation_method") or ""), item.get("depreciation_method")),
                    item.get("original_value"),
                    item.get("impairment_value"),
                    item.get("depreciated_periods"),
                    item.get("accumulated_depreciation"),
                    item.get("residual_rate"),
                    item.get("accumulated_depreciation_account_code"),
                    item.get("expense_account_code"),
                    item.get("asset_account_code"),
                ]
            )
        stream = BytesIO()
        wb.save(stream)
        stream.seek(0)
        return stream

    async def import_excel(self, tenant_id: int, content: bytes, user: User) -> dict[str, Any]:
        from apps.kuaicaiwu.services.fa_asset_import_template import (
            resolve_header_index_map,
            row_to_create_payload,
        )

        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {"created": 0, "errors": ["文件为空"]}
        header_row_idx = 0
        for idx, row in enumerate(rows[:5]):
            if row and any(str(cell or "").strip() for cell in row):
                header_row_idx = idx
                break
        index_map = resolve_header_index_map(list(rows[header_row_idx]))
        if "asset_name" not in index_map:
            raise ValidationError("未识别导入表头，请使用系统提供的固定资产导入模板")
        created = 0
        errors: list[str] = []
        for idx, row in enumerate(rows[header_row_idx + 1 :], start=header_row_idx + 2):
            if not row or not any(row):
                continue
            payload = row_to_create_payload(tuple(row), index_map)
            if not payload.get("asset_name"):
                errors.append(f"第{idx}行：资产名称不能为空")
                continue
            if not payload.get("category_name"):
                errors.append(f"第{idx}行：资产类别不能为空")
                continue
            payload["change_method"] = payload.get("change_method") or "import"
            try:
                await self.create_asset(tenant_id, payload, user)
                created += 1
            except Exception as exc:
                errors.append(f"第{idx}行：{exc}")
        return {"created": created, "errors": errors}
